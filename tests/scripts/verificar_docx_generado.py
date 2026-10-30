from docx import Document
from pathlib import Path
import os
import re

p=Path(os.environ.get("CAMILA_PPA_QA_OUTPUT","artifacts/camila-ppa-6-semanas-ficticio.docx"))
doc=Document(p)
assert len(doc.tables)==7, f"Esperadas 7 tablas, recibidas {len(doc.tables)}"
tabla=doc.tables[4]
assert len(tabla.rows)==31, f"Se requieren 1 cabecera + 6 títulos semanales + 24 filas de áreas: {len(tabla.rows)}"
titulos=[r.cells[0].text for r in tabla.rows if r.cells[0].text.startswith("SEMANA ")]
assert len(titulos)==6, titulos
for i,titulo in enumerate(titulos,1):
    assert titulo.startswith(f"SEMANA {i} "), titulo
total="\\n".join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for r in t.rows for c in r.cells])
footer="\\n".join(p.text for sec in doc.sections for p in sec.footer.paragraphs)
assert "Docente Ficticia Nueva" in total+footer
assert "Docente Ficticia Antigua" not in total+footer
assert "Institución Ficticia Nueva" in total
assert "Institución Ficticia Anterior" not in total
assert "26-10-26" not in total+footer
assert "TÍTULO ANTIGUO" not in total+footer
assert "10-12-26" in total
print("OK: Word abre; 7 tablas; 6 semanas; 24 áreas; identidad nueva; pie actualizado; sin fechas ajenas")
