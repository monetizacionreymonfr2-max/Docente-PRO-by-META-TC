"""Genera un Word institucional completamente ficticio para pruebas visuales locales."""

from io import BytesIO
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


def poner_bordes(tabla):
    propiedades = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "single")
        borde.set(qn("w:sz"), "5")
        borde.set(qn("w:color"), "D9D9D9")
        bordes.append(borde)
    propiedades.append(bordes)


def sombrear(celda, color):
    propiedades = celda._tc.get_or_add_tcPr()
    fondo = OxmlElement("w:shd")
    fondo.set(qn("w:fill"), color)
    propiedades.append(fondo)


def preparar_logo():
    imagen = Image.new("RGB", (180, 180), "white")
    dibujo = ImageDraw.Draw(imagen)
    dibujo.ellipse((13, 13, 167, 167), fill="#1E3A8A", outline="#D9D9D9", width=4)
    try:
        fuente = ImageFont.truetype("DejaVuSans-Bold.ttf", 104)
    except OSError:
        fuente = ImageFont.load_default()
    caja = dibujo.textbbox((0, 0), "F", font=fuente)
    dibujo.text(((180 - (caja[2] - caja[0])) / 2 - caja[0],
                 (180 - (caja[3] - caja[1])) / 2 - caja[1]),
                "F", fill="white", font=fuente)
    buffer = BytesIO()
    imagen.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def generar(destino):
    documento = Document()
    seccion = documento.sections[0]
    seccion.page_width = Inches(8.5)
    seccion.page_height = Inches(11)
    seccion.top_margin = Inches(0.72)
    seccion.bottom_margin = Inches(0.72)
    seccion.left_margin = Inches(0.72)
    seccion.right_margin = Inches(0.72)

    normal = documento.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    encabezado = seccion.header.paragraphs[0]
    encabezado.alignment = WD_ALIGN_PARAGRAPH.CENTER
    encabezado.add_run().add_picture(preparar_logo(), width=Inches(0.38))
    encabezado.add_run("  ESCUELA FICTICIA DE PRUEBAS").bold = True

    titulo = documento.add_paragraph(style="Title")
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo_run = titulo.add_run("Planificación de Aprendizaje de Prueba")
    titulo_run.font.color.rgb = RGBColor(0, 0, 0)
    titulo_run.font.size = Pt(16)
    # Word Title trae en algunas plantillas una regla azul heredada.
    for propiedades in (documento.styles["Title"]._element.pPr, titulo._p.get_or_add_pPr()):
        if propiedades is not None:
            for borde in propiedades.findall(qn("w:pBdr")):
                propiedades.remove(borde)

    descripcion = documento.add_paragraph(
        "Documento de simulación local con información inventada para comprobar "
        "la integridad de Mi Formato y la legibilidad de su vista previa."
    )
    descripcion.paragraph_format.space_after = Pt(10)

    datos = documento.add_table(rows=3, cols=2)
    datos.autofit = False
    for fila, (etiqueta, valor) in zip(datos.rows, [
        ("Docente", "Docente Ficticio A"),
        ("Grado y sección", "1.er grado A"),
        ("Proyecto", "Exploramos palabras y convivencia"),
    ]):
        fila.cells[0].text = etiqueta
        fila.cells[1].text = valor
        fila.cells[0].width = Inches(1.8)
        fila.cells[1].width = Inches(5.0)
        fila.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        fila.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for celda in fila.cells:
            celda.paragraphs[0].paragraph_format.space_after = Pt(3)
    poner_bordes(datos)

    documento.add_paragraph("Cronograma ficticio", style="Heading 2").runs[0].font.color.rgb = RGBColor(0, 0, 0)
    cronograma = documento.add_table(rows=1, cols=3)
    cronograma.autofit = False
    cronograma.columns[0].width = Inches(1.1)
    cronograma.columns[1].width = Inches(1.5)
    cronograma.columns[2].width = Inches(4.1)
    for celda, titulo_columna in zip(cronograma.rows[0].cells, ("Semana", "Área", "Actividad")):
        celda.text = titulo_columna
        sombrear(celda, "1E3A8A")
        for run in celda.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    propiedades_encabezado = cronograma.rows[0]._tr.get_or_add_trPr()
    propiedades_encabezado.append(OxmlElement("w:tblHeader"))
    for semana, area, actividad in [
        ("1", "Lengua", "Lectura compartida de un relato inventado."),
        ("2", "Matemática", "Conteo lúdico con objetos de cartón."),
        ("3", "Ciencias", "Observación de una planta de papel."),
    ]:
        fila = cronograma.add_row()
        for celda, texto in zip(fila.cells, (semana, area, actividad)):
            celda.text = texto
            celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    poner_bordes(cronograma)

    documento.add_paragraph("Evaluación ficticia", style="Heading 2").runs[0].font.color.rgb = RGBColor(0, 0, 0)
    evaluacion = documento.add_table(rows=2, cols=2)
    evaluacion.cell(0, 0).merge(evaluacion.cell(0, 1)).text = "Criterios de evaluación"
    evaluacion.cell(1, 0).text = "Participa en lectura compartida"
    evaluacion.cell(1, 1).text = "Lista de cotejo"
    sombrear(evaluacion.cell(0, 0), "E8EFF9")
    poner_bordes(evaluacion)

    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.add_run("Material ficticio para verificación local de DOCENTY PRO").italic = True

    archivo = Path(destino)
    archivo.parent.mkdir(parents=True, exist_ok=True)
    documento.save(archivo)
    print(f"Plantilla de prueba creada: {archivo} ({archivo.stat().st_size} bytes)")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Uso: generar_word_visual_ficticio.py <ruta-salida.docx>")
    generar(sys.argv[1])
