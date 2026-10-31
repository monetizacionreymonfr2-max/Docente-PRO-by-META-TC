import io
import os
import unittest
import zipfile
from pathlib import Path
import sys
sys.path.insert(0, str(Path(__file__).parent.parent / "workers" / "libreoffice"))
from server import signature, check_request, validate_docx, convert_docx

KEY = "a-secret-for-testing-at-least-32-character-long"

def tiny_docx():
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,"w") as z:
        z.writestr("[Content_Types].xml","<Types/>")
        z.writestr("word/document.xml","<w:document/>")
    return stream.getvalue()

class LibreOfficeWorkerTests(unittest.TestCase):
    def test_hmac_ok(self):
        payload=tiny_docx()
        self.assertTrue(check_request(KEY,"1000",signature(KEY,"1000",payload),payload,now=1000))
    def test_reject_replay_and_tampering(self):
        payload=tiny_docx()
        sig=signature(KEY,"1000",payload)
        self.assertFalse(check_request(KEY,"1000",sig,payload+b"x",now=1000))
        self.assertFalse(check_request(KEY,"1000",sig,payload,now=1130))
        self.assertFalse(check_request(KEY,"not-time",sig,payload,now=1000))
    def test_no_zip_bombs_or_macros(self):
        data=io.BytesIO()
        with zipfile.ZipFile(data,"w") as z:
            z.writestr("[Content_Types].xml","ok")
            z.writestr("word/document.xml","ok")
            z.writestr("word/vbaProject.bin","malware")
        with self.assertRaises(ValueError):
            validate_docx(data.getvalue())
        with self.assertRaises(ValueError):
            validate_docx(b"not-a-word")
    @unittest.skipUnless(os.path.isfile("/usr/bin/soffice"),"LibreOffice not installed")
    def test_real_docx_to_pdf(self):
        from docx import Document
        d=Document()
        d.add_heading("PPA ficticio de prueba",0)
        table=d.add_table(rows=2,cols=2)
        table.cell(0,0).text="Grado"
        table.cell(0,1).text="Segundo"
        table.cell(1,0).text="Semanas"
        table.cell(1,1).text="Seis"
        blob=io.BytesIO();d.save(blob)
        result=convert_docx(blob.getvalue())
        self.assertTrue(result.startswith(b"%PDF-"))
        self.assertGreater(len(result),500)

if __name__=="__main__":
    unittest.main()
